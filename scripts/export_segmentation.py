"""
Полная выгрузка гибридного разбора документов (docpipe, режим hybrid) — для проверки человеком
и подбора порогов: каждый фрагмент, кто из алгоритмов голосовал за его границу, ключевые слова,
все кандидаты-подэтапы со сходством и кто выбрал итог (алгоритм или модель).

Запуск на сервере:  python scripts/export_segmentation.py [имя_файла.docx ...]
Результат:          Разбор_фрагментов.xlsx (без аргументов — все документы)
"""
import sys

import _path  # noqa: F401 — backend/ в sys.path
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import db
import docpipe.store as store

OUT = "Разбор_фрагментов.xlsx"
COLUMNS = ("Документ", "Раздел", "№ фрагмента", "Текст фрагмента", "TextTiling", "C99", "Разметка текста",
           "Вес границы", "Ключевые слова", "Кандидаты (сходство)", "Итог: подэтапы", "Кто выбрал",
           "Общая информация")
LABELS = {"lexical": "алгоритм (без ИИ)", "local": "локальная модель", "deepseek": "DeepSeek"}


def rows(filenames):
    where, params = "", ()
    if filenames:
        where, params = "WHERE d.filename = ANY(%s)", (list(filenames),)
    return db.query(
        "SELECT d.filename, s.seq AS section_seq, s.heading_path, c.seq, c.text, c.substages, "
        "       c.is_general, c.meta "
        f"FROM chunks c JOIN sections s ON s.id = c.section_id JOIN documents d ON d.id = s.doc_id {where} "
        "ORDER BY d.filename, s.seq, c.seq", params) or []


def main(argv):
    titles = {sub["id"]: f"{st.get('title', '')} / {sub.get('title', '')}"
              for st in store.get_plan_structure().get("stages") or [] for sub in st.get("substages") or []}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Фрагменты"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="3B5B8C")
    n = 0
    for r in rows(argv):
        meta = r.get("meta") or {}
        votes = meta.get("votes") or {}
        mark = lambda k: "✓" if votes.get(k) else ""                                # noqa: E731
        ws.append((r["filename"], " / ".join(r.get("heading_path") or []), f"{r['section_seq'] + 1}.{r['seq'] + 1}",
                   r["text"], mark("texttiling"), mark("c99"), mark("cues"), meta.get("boundary_score"),
                   ", ".join(w for w, _ in meta.get("keywords") or []),
                   "\n".join(f"{c['title']} — {c['score']}" for c in meta.get("candidates") or []),
                   "\n".join(titles.get(s["id"], s["id"]) for s in r.get("substages") or []),
                   LABELS.get(meta.get("picked_by"), meta.get("picked_by") or ""),
                   "да" if r.get("is_general") else ""))
        n += 1
    for col, width in zip("ABCDEFGHIJKLM", (28, 30, 10, 80, 10, 8, 10, 10, 40, 50, 45, 18, 10)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(OUT)
    print(f"Фрагментов: {n} -> {OUT}" + ("" if n else " (нет фрагментов: документы размечены в режиме llm "
                                                    "или ещё не разобраны)"))


if __name__ == "__main__":
    main(sys.argv[1:])
